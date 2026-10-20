"""Export the whole recipe database as an .xlsx workbook.

Built for dropping into a Claude project folder (or any spreadsheet app):
one self-contained row per recipe on the Recipes sheet, plus long-format
Ingredients, Tags and (when anny_food has filled it) per-serving Nutrition
sheets, and a README sheet that explains the columns.

Read-only: opens its own connection, never writes. Quantities follow the
View dialog's rule — qty_per_serving × base_servings, falling back to the
verbatim qty_raw only when there is no number.
"""
from __future__ import annotations

import io
import json
from collections import defaultdict
from datetime import date
from pathlib import Path

from meal_planner import db as _db

_FONT = "Arial"
_HEADER_FILL = "4F6D7A"


def export_filename(today: date | None = None) -> str:
    return f"recipes-export-{(today or date.today()).isoformat()}.xlsx"


def _fmt_qty(qty: float) -> str:
    qty = round(qty, 3)
    return str(int(qty)) if qty == int(qty) else f"{qty:g}"


def _qty_for_servings(ing, servings: int) -> str:
    if ing["qty_per_serving"] is not None:
        return _fmt_qty(ing["qty_per_serving"] * servings)
    return ing["qty_raw"] or ""


def _ingredient_line(ing, servings: int) -> str:
    parts = [p for p in (_qty_for_servings(ing, servings), ing["unit"], ing["name"]) if p]
    line = " ".join(parts)
    if ing["notes"]:
        line += f" ({ing['notes']})"
    return "• " + line


def _load(conn) -> dict:
    recipes = conn.execute(
        "SELECT * FROM recipes ORDER BY title COLLATE NOCASE, id"
    ).fetchall()
    ingredients = defaultdict(list)
    for row in conn.execute("SELECT * FROM ingredients ORDER BY recipe_id, sort_order, id"):
        ingredients[row["recipe_id"]].append(row)
    tags_by_recipe = defaultdict(list)
    for row in conn.execute(
        "SELECT rt.recipe_id, t.name FROM recipe_tags rt "
        "JOIN tags t ON t.id = rt.tag_id ORDER BY t.name COLLATE NOCASE"
    ):
        tags_by_recipe[row["recipe_id"]].append(row["name"])
    tags = [r["name"] for r in conn.execute("SELECT name FROM tags ORDER BY name COLLATE NOCASE")]
    # anny_recipe_nutrition belongs to anny_food and may not exist yet.
    has_nutrition = conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = 'anny_recipe_nutrition'"
    ).fetchone()
    nutrition = (
        conn.execute("SELECT * FROM anny_recipe_nutrition").fetchall() if has_nutrition else []
    )
    return {
        "recipes": recipes,
        "ingredients": ingredients,
        "tags_by_recipe": tags_by_recipe,
        "tags": tags,
        "nutrition": nutrition,
    }


def build_workbook(*, path: Path | None = None, today: date | None = None):
    """Return an openpyxl Workbook of the full recipe database at ``path``."""
    # Imported here so a console venv without openpyxl still renders the tab.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    with _db._get_conn(path or _db.DB_PATH) as conn:
        data = _load(conn)
    recipes = data["recipes"]
    ingredients = data["ingredients"]
    tags_by_recipe = data["tags_by_recipe"]
    title_by_id = {r["id"]: r["title"] for r in recipes}

    header_font = Font(name=_FONT, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=_HEADER_FILL)
    body_font = Font(name=_FONT, size=10)
    wrap_top = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")

    def write_table(ws, headers, rows, widths, wrap_cols=()):
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill = header_font, header_fill
        for row in rows:
            ws.append(row)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = wrap_top if cell.column in wrap_cols else top
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions

    wb = Workbook()

    ws = wb.active
    ws.title = "Recipes"
    write_table(
        ws,
        ["Recipe ID", "Title", "Servings", "Cook time (min)", "Tags", "Recipe book", "Source",
         "# Ingredients", "Ingredients (for listed servings)", "Instructions", "Created", "Updated"],
        [
            [
                r["id"], r["title"], r["base_servings"], r["cook_time_min"] or None,
                ", ".join(tags_by_recipe[r["id"]]), r["recipe_book"] or "", r["source"] or "",
                len(ingredients[r["id"]]),
                "\n".join(_ingredient_line(i, r["base_servings"]) for i in ingredients[r["id"]]),
                (r["instructions"] or "").strip(),
                r["created_at"][:10], r["updated_at"][:10],
            ]
            for r in recipes
        ],
        [9, 38, 9, 10, 30, 22, 30, 8, 55, 80, 11, 11],
        wrap_cols={2, 5, 6, 7, 9, 10},
    )

    ws = wb.create_sheet("Ingredients")
    write_table(
        ws,
        ["Recipe ID", "Recipe", "Order", "Ingredient", "Qty (for servings)", "Unit",
         "Qty per serving", "Servings", "Notes", "Grocery section"],
        [
            [
                r["id"], r["title"], i["sort_order"], i["name"],
                _qty_for_servings(i, r["base_servings"]), i["unit"] or "",
                round(i["qty_per_serving"], 4) if i["qty_per_serving"] is not None else None,
                r["base_servings"], i["notes"] or "", i["todoist_section"] or "",
            ]
            for r in recipes
            for i in ingredients[r["id"]]
        ],
        [9, 34, 7, 34, 12, 12, 11, 9, 40, 18],
        wrap_cols={2, 4, 9},
    )

    recipes_by_tag = defaultdict(list)
    for r in recipes:
        for name in tags_by_recipe[r["id"]]:
            recipes_by_tag[name].append(r["title"])
    ws = wb.create_sheet("Tags")
    write_table(
        ws,
        ["Tag", "# Recipes", "Recipes"],
        [[t, len(recipes_by_tag[t]), "; ".join(recipes_by_tag[t])] for t in data["tags"]],
        [24, 10, 120],
        wrap_cols={3},
    )

    nutrition = [n for n in data["nutrition"] if n["recipe_id"] in title_by_id]
    nutrition.sort(key=lambda n: title_by_id[n["recipe_id"]].lower())
    if nutrition:
        keys = sorted({k for n in nutrition for k in json.loads(n["values_json"])})
        ws = wb.create_sheet("Nutrition (per serving)")
        rows = []
        for n in nutrition:
            values = json.loads(n["values_json"])
            rows.append([
                n["recipe_id"], title_by_id[n["recipe_id"]], *[values.get(k) for k in keys],
                n["confidence"], ", ".join(json.loads(n["safety_flags_json"] or "[]")),
                n["notes"] or "",
            ])
        write_table(
            ws,
            ["Recipe ID", "Recipe", *keys, "Confidence", "Safety flags", "Notes"],
            rows,
            [9, 34, *[12] * len(keys), 11, 22, 60],
            wrap_cols={2, len(keys) + 5},
        )

    n_ingredients = sum(len(v) for v in ingredients.values())
    readme = [
        ("Household recipe database export", "title"),
        (f"Exported {(today or date.today()).isoformat()} from the Mini Ops console Recipes tab "
         "(http://homeserver:8503/?tab=recipes).", None),
        (f"{len(recipes)} recipes · {n_ingredients} ingredient rows · {len(data['tags'])} tags"
         + (f" · {len(nutrition)} nutrition rows" if nutrition else ""), None),
        ("", None),
        ("Sheets", "heading"),
        ("Recipes — one row per recipe, self-contained: servings, tags, full ingredient list "
         "and instructions.", None),
        ("Ingredients — one row per ingredient (join to Recipes on Recipe ID). Good for "
         "grocery/pantry analysis.", None),
        ("Tags — each tag with the recipes that carry it.", None),
    ]
    if nutrition:
        readme.append((
            "Nutrition (per serving) — Claude-estimated per-serving food-group amounts "
            "(pregnancy tracker); estimates, not lab values.", None,
        ))
    readme += [
        ("", None),
        ("Notes", "heading"),
        ("Ingredient quantities are for the recipe's Servings value. To scale: Qty per serving "
         "× desired servings. Blank quantity = to taste / unspecified.", None),
        ("Grocery section = the Todoist grocery list section the item is sent to "
         "('Skip' = pantry staple, not sent).", None),
        ("Blank Cook time / Source / Instructions = not recorded in the database.", None),
        ("This is a read-only snapshot; the recipe database on the homeserver remains the "
         "source of truth.", None),
    ]
    ws = wb.create_sheet("README", 0)
    for text, kind in readme:
        ws.append([text])
        size = 14 if kind == "title" else 10
        ws.cell(row=ws.max_row, column=1).font = Font(name=_FONT, size=size, bold=kind is not None)
    ws.column_dimensions["A"].width = 120
    wb.active = 0
    return wb


def build_workbook_bytes(*, path: Path | None = None, today: date | None = None) -> bytes:
    buf = io.BytesIO()
    build_workbook(path=path, today=today).save(buf)
    return buf.getvalue()
