"""Tests for meal_planner.export_xlsx (Recipes tab "Export to Excel")."""
from __future__ import annotations

import io
import json
import sqlite3
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from anny_food import db as anny_db
from meal_planner.db import add_recipe_tag, init_db, insert_ingredient, insert_recipe
from meal_planner.export_xlsx import build_workbook_bytes, export_filename


@pytest.fixture
def db_path(tmp_path: Path) -> Path:
    p = tmp_path / "recipes.db"
    init_db(p)
    soup = insert_recipe(
        title="Soup", base_servings=4, instructions="1. boil\n2. eat",
        cook_time_min=30, source="nas-intake", recipe_book="BP recipes", path=p,
    )
    insert_ingredient(recipe_id=soup, name="carrots", qty_per_serving=0.5, unit="cup",
                      notes="diced", todoist_section="Fruits + Veggies", sort_order=0, path=p)
    insert_ingredient(recipe_id=soup, name="salt", sort_order=1, path=p)
    add_recipe_tag(soup, "vegetarian", path=p)
    add_recipe_tag(soup, "dinner", path=p)
    apple = insert_recipe(title="apple crisp", base_servings=6, path=p)
    add_recipe_tag(apple, "dinner", path=p)
    return p


def _load(p: Path, **kw):
    return load_workbook(io.BytesIO(build_workbook_bytes(path=p, **kw)))


def _rows(ws) -> list[dict]:
    header, *body = ws.iter_rows(values_only=True)
    return [dict(zip(header, row)) for row in body]


def test_sheets_and_readme_first(db_path: Path) -> None:
    wb = _load(db_path, today=date(2026, 10, 20))
    assert wb.sheetnames == ["README", "Recipes", "Ingredients", "Tags"]
    readme = [r[0] for r in wb["README"].iter_rows(values_only=True)]
    assert "Exported 2026-10-20" in readme[1]
    assert readme[2].startswith("2 recipes · 2 ingredient rows · 2 tags")


def test_recipes_sheet_is_self_contained_and_case_insensitively_sorted(db_path: Path) -> None:
    rows = _rows(_load(db_path)["Recipes"])
    assert [r["Title"] for r in rows] == ["apple crisp", "Soup"]
    soup = rows[1]
    assert soup["Servings"] == 4
    assert soup["Cook time (min)"] == 30
    assert soup["Tags"] == "dinner, vegetarian"
    assert soup["Recipe book"] == "BP recipes"
    assert soup["# Ingredients"] == 2
    assert soup["Ingredients (for listed servings)"] == "• 2 cup carrots (diced)\n• salt"
    assert soup["Instructions"] == "1. boil\n2. eat"
    assert rows[0]["Cook time (min)"] is None
    assert rows[0]["# Ingredients"] == 0


def test_ingredients_sheet_long_format(db_path: Path) -> None:
    rows = _rows(_load(db_path)["Ingredients"])
    assert [r["Ingredient"] for r in rows] == ["carrots", "salt"]
    carrots = rows[0]
    assert carrots["Recipe"] == "Soup"
    assert carrots["Qty (for servings)"] == "2"
    assert carrots["Qty per serving"] == 0.5
    assert carrots["Grocery section"] == "Fruits + Veggies"
    assert rows[1]["Qty (for servings)"] in ("", None)


def test_quantity_prefers_scaled_number_over_stale_qty_raw(db_path: Path) -> None:
    # qty_raw is the text captured at insert; an edit updates only
    # qty_per_serving, so the scaled number wins (same rule as the View dialog).
    with sqlite3.connect(db_path) as conn:
        conn.execute("UPDATE ingredients SET qty_raw = '5' WHERE name = 'carrots'")
        conn.execute("UPDATE ingredients SET qty_raw = '2-3' WHERE name = 'salt'")
    rows = _rows(_load(db_path)["Ingredients"])
    assert rows[0]["Qty (for servings)"] == "2"
    assert rows[1]["Qty (for servings)"] == "2-3"


def test_fractional_quantities_are_rounded(db_path: Path) -> None:
    with sqlite3.connect(db_path) as conn:
        conn.execute("UPDATE ingredients SET qty_per_serving = 1.0/3 WHERE name = 'carrots'")
    rows = _rows(_load(db_path)["Ingredients"])
    assert rows[0]["Qty (for servings)"] == "1.333"


def test_tags_sheet_lists_recipes(db_path: Path) -> None:
    rows = {r["Tag"]: r for r in _rows(_load(db_path)["Tags"])}
    assert rows["dinner"]["# Recipes"] == 2
    assert rows["dinner"]["Recipes"] == "apple crisp; Soup"
    assert rows["vegetarian"]["Recipes"] == "Soup"


def test_nutrition_sheet_when_anny_food_has_rows(db_path: Path) -> None:
    anny_db.init_db(db_path)
    with sqlite3.connect(db_path) as conn:
        soup_id = conn.execute("SELECT id FROM recipes WHERE title = 'Soup'").fetchone()[0]
        conn.execute(
            "INSERT INTO anny_recipe_nutrition (recipe_id, ingredients_hash, values_json, "
            "breakdown_json, safety_flags_json, confidence, notes, model, created_at) "
            "VALUES (?, 'h', ?, '[]', ?, 'med', 'rough', 'claude', '2026-10-20')",
            (soup_id, json.dumps({"veg_other": 0.5, "calcium": 20}), json.dumps(["raw_egg"])),
        )
    wb = _load(db_path)
    assert wb.sheetnames[-1] == "Nutrition (per serving)"
    (row,) = _rows(wb["Nutrition (per serving)"])
    assert row["Recipe"] == "Soup"
    assert row["calcium"] == 20
    assert row["veg_other"] == 0.5
    assert row["Confidence"] == "med"
    assert row["Safety flags"] == "raw_egg"


def test_export_does_not_write(db_path: Path) -> None:
    before = db_path.read_bytes()
    build_workbook_bytes(path=db_path)
    assert db_path.read_bytes() == before


def test_export_filename() -> None:
    assert export_filename(date(2026, 10, 20)) == "recipes-export-2026-10-20.xlsx"
