"""The Recipes tab renders the "Export to Excel" download button (real Streamlit)."""
from __future__ import annotations

import io
import textwrap
from pathlib import Path

from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from console.tabs import plan
from meal_planner import db as _db
from meal_planner.db import init_db, insert_recipe


def _seed(tmp_path: Path, monkeypatch) -> Path:
    p = tmp_path / "recipes.db"
    init_db(p)
    monkeypatch.setattr(_db, "DB_PATH", p)
    insert_recipe(title="Soup", base_servings=4, path=p)
    return p


def test_recipes_tab_renders_export_button(tmp_path: Path, monkeypatch) -> None:
    _seed(tmp_path, monkeypatch)
    script = tmp_path / "recipes_tab_app.py"
    script.write_text(textwrap.dedent("""
        from console.tabs import plan
        plan.render()
    """))
    at = AppTest.from_file(str(script), default_timeout=30)
    at.run()
    assert not at.exception
    assert not at.error
    (button,) = at.get("download_button")
    assert button.proto.label == "Export to Excel"
    assert button.proto.help.endswith("Ignores the filters above.")


def test_build_export_returns_whole_db(tmp_path: Path, monkeypatch) -> None:
    _seed(tmp_path, monkeypatch)
    wb = load_workbook(io.BytesIO(plan._build_export()))
    titles = [r[1] for r in wb["Recipes"].iter_rows(min_row=2, values_only=True)]
    assert titles == ["Soup"]
